"""
Phase 6 — Document Generation
Fills Individual Template.docx with Phase 1 + Phase 5 output.
Saves to outputs/<PersonName>/<PersonName>.docx — folder replaced on every run.
"""
import re
import os
import sys
import io
import shutil
import datetime
from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.shared import Inches

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.dirname(__file__))
from config import TEMPLATE_PATH, OUTPUTS_DIR
from document_utils import populate_document_placeholders_with_dates


# ── FORMAT PHASE 1 ────────────────────────────────────────────────────────────

def format_phase1(p1: dict) -> dict:
    age_raw = p1.get("age")
    age_src = p1.get("age_source", "")
    age_val = str(age_raw) if age_raw is not None else "Not found"

    nat_val = str(p1.get("nationality") or "Not found")
    nat_src = p1.get("nat_source", "")

    nw_raw = p1.get("net_worth")
    nw_src = p1.get("nw_source", "")
    nw_val = str(nw_raw) if nw_raw else "Not found"

    degrees = p1.get("education", [])
    edu_src = p1.get("edu_source", "")
    if isinstance(degrees, list) and degrees:
        edu_val = "; ".join(
            f"{d.get('degree', '')} ({d.get('institution', '')})"
            for d in degrees if isinstance(d, dict) and d.get("degree")
        )
    else:
        edu_val = "Not found"

    cls_yr   = p1.get("class_of_year") or p1.get("graduation_year")
    cls_inst = p1.get("class_of_institution") or p1.get("institution", "")
    cls_val  = f"{cls_yr}" + (f", {cls_inst}" if cls_inst else "") if cls_yr else "Not found"

    return {
        "NAME":             p1.get("name", ""),
        "AGE":              age_val,
        "AGE_SRC":          age_src,
        "NATIONALITY":      nat_val,
        "NAT_SRC":          nat_src,
        "NET_WORTH":        nw_val,
        "NW_SRC":           nw_src,
        "EDUCATION_DEGREE": edu_val,
        "EDU_SRC":          edu_src,
        "CLASS_OF":         cls_val,
    }


# ── FORMAT PHASE 5 ────────────────────────────────────────────────────────────

def format_phase5(p5: dict) -> dict:
    results = p5.get("results", {})
    fields  = {}

    attr_to_placeholder = {
        "BIOGRAPHY":            "BIOGRAPHY",
        "FAMILY":               "FAMILY",
        "INTERESTING_FACTS":    "INTERESTING_FACTS",
        "ADVERSE_NEWS":         "ADVERSE_NEWS",
        "GIVING":               "DONATION_HISTORY",
        "POTENTIAL_CONNECTORS": "POTENTIAL_CONNECTORS",
        "KEY_POSITIONS":        "KEY_POSITION_ORGANISATION",
    }

    for attr, placeholder in attr_to_placeholder.items():
        data    = results.get(attr, {})
        content = data.get("content") or "No information found."
        sources = data.get("sources", [])
        # Sources can be plain URL strings or dicts with url/date keys
        source_tuples = []
        for s in sources:
            if isinstance(s, dict):
                source_tuples.append((s.get("url", ""), s.get("date", "")))
            elif isinstance(s, str) and s:
                source_tuples.append((s, ""))
        fields[placeholder] = [(content, source_tuples)]

    key_data = results.get("KEY_POSITIONS", {})
    fields["JOB_TITLE"] = [(key_data.get("content") or "Not found", [])]
    # DEMONSTRATED_INTERESTS has its own section in results
    demo_data = results.get("DEMONSTRATED_INTERESTS", {})
    demo_content = demo_data.get("content") or "No information found."
    demo_sources = demo_data.get("sources", [])
    demo_tuples = []
    for s in demo_sources:
        if isinstance(s, dict):
            demo_tuples.append((s.get("url", ""), s.get("date", "")))
        elif isinstance(s, str) and s:
            demo_tuples.append((s, ""))
    fields["DEMONSTRATED_INTERESTS"] = [(demo_content, demo_tuples)]
    return fields


# ── INSERT PHOTO directly via python-docx ────────────────────────────────────

def insert_photo(doc: Document, photo_path: str):
    """
    Replace [[IMAGE]] placeholder with actual photo using python-docx only.
    Iterates ALL paragraphs including inside table cells via XML iterator.
    """
    if not photo_path or not Path(photo_path).exists():
        print(f"  ⚠️  No photo to insert")
        return

    tag     = "[[IMAGE]]"
    pattern = re.compile(re.escape(tag))

    # Prepare image bytes
    try:
        from PIL import Image as PILImage
        with open(photo_path, "rb") as f:
            img_bytes = io.BytesIO(f.read())
        pil_img = PILImage.open(img_bytes)
        if pil_img.mode != "RGB":
            pil_img = pil_img.convert("RGB")
        converted = io.BytesIO()
        pil_img.save(converted, format="PNG")
        converted.seek(0)
        image_stream = converted
    except ImportError:
        # Pillow not installed — use raw bytes
        image_stream = photo_path

    # Iterate ALL paragraphs in entire document XML
    for para_element in doc.element.body.iter(qn("w:p")):
        para      = Paragraph(para_element, para_element.getparent())
        full_text = "".join(run.text for run in para.runs)
        if not pattern.search(full_text):
            continue

        # Found [[IMAGE]] — clear all runs and insert image
        p = para._p
        # Remove all w:r (run) elements
        for r_elem in p.findall(qn("w:r")):
            p.remove(r_elem)
        # Remove bookmarks that may interfere
        for bk in p.findall(qn("w:bookmarkStart")):
            p.remove(bk)
        for bk in p.findall(qn("w:bookmarkEnd")):
            p.remove(bk)

        # Add image using document part directly (works inside table cells)
        if isinstance(image_stream, str):
            with open(image_stream, "rb") as f:
                img_data = io.BytesIO(f.read())
        else:
            image_stream.seek(0)
            img_data = image_stream

        # Get the document part via the top-level document object
        from docx.oxml import OxmlElement
        from docx.opc.part import Part
        pic_part = doc.part
        img_data.seek(0)
        r_elem = OxmlElement("w:r")
        drawing = pic_part.new_pic_inline(img_data, width=Inches(1.5), height=None)
        # Create run with drawing
        run_xml = OxmlElement("w:r")
        drawing_xml = OxmlElement("w:drawing")
        drawing_xml.append(drawing)
        run_xml.append(drawing_xml)
        p.append(run_xml)

        print(f"  ✅ Photo inserted into document")
        return

    print(f"  ⚠️  [[IMAGE]] placeholder not found in document")


# ── MAIN ──────────────────────────────────────────────────────────────────────

def run(phase1_output: dict, phase5_output: dict, photo_path: str = None) -> str:
    name = phase5_output.get("name", phase1_output.get("name", "Unknown"))

    print(f"\n  Loading template...")
    doc = Document(str(TEMPLATE_PATH))

    # Fill text placeholders
    p1_fields = format_phase1(phase1_output)
    p5_fields = format_phase5(phase5_output)
    all_fields = {}
    src_keys = {"AGE_SRC", "NAT_SRC", "NW_SRC", "EDU_SRC"}

    # Get any available source URL as fallback for all shallow fields
    fallback_src = (p1_fields.get("AGE_SRC") or p1_fields.get("NAT_SRC") or
                    p1_fields.get("NW_SRC") or p1_fields.get("EDU_SRC") or "")

    for k, v in p1_fields.items():
        if k in src_keys:
            continue  # skip raw source keys
        src_url = p1_fields.get(k + "_SRC", "") or fallback_src
        if src_url and v and v != "Not found" and k != "NAME":
            all_fields[k] = [(f"{v} [[source]]", [(src_url, "")])]
        else:
            all_fields[k] = [(v, [])]
    all_fields.update(p5_fields)

    # Insert photo FIRST before populate consumes [[IMAGE]]
    insert_photo(doc, photo_path)

    print(f"  Filling {len(all_fields)} placeholders...")
    populate_document_placeholders_with_dates(doc, all_fields)

    # Save to outputs/<PersonName>/ — replace folder on every run
    safe_name  = re.sub(r"[^\w\s-]", "", name).strip().replace(" ", "_")
    person_dir = OUTPUTS_DIR / safe_name
    if person_dir.exists():
        import time
        # On Windows files may be locked — try multiple approaches
        for attempt in range(5):
            try:
                shutil.rmtree(person_dir, ignore_errors=True)
                if not person_dir.exists():
                    break
            except Exception:
                pass
            time.sleep(1)
    person_dir.mkdir(parents=True, exist_ok=True)

    output_path = str(person_dir / f"{safe_name}.docx")
    doc.save(output_path)
    print(f"  ✅ Document saved: {output_path}")
    return output_path
