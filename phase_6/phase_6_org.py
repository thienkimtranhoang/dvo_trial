"""
Phase 6 — Organisation Document Generation
Fills Organisation Template.docx with Phase 4 output.
Saves to outputs/<OrgName>/<OrgName>.docx
"""
import re
import os
import sys
import shutil
from pathlib import Path
from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.dirname(__file__))
from config import OUTPUTS_DIR
from document_utils import populate_document_placeholders_with_dates

TEMPLATE_PATH = Path(__file__).parent / "template" / "Organisation Template.docx"


def format_org_fields(phase4_results: dict, name: str) -> dict:
    """Convert phase4 results to field format for document_utils."""
    fields = {}

    attr_to_placeholder = {
        "DATE_OF_ESTABLISHMENT":  "DATE_OF_ESTABLISHMENT",
        "BIOGRAPHY":              "BIOGRAPHY",
        "GIVING":                 "DONATION_HISTORY",
        "DEMONSTRATED_INTERESTS": "DEMONSTRATED_INTERESTS",
        "OTHER_INTERESTING_FACTS":"OTHER_INTERESTING_FACTS",
        "POTENTIAL_CONNECTORS":   "POTENTIAL_CONNECTORS",
        "ADVERSE_NEWS":           "ADVERSE_NEWS",
    }

    for attr, placeholder in attr_to_placeholder.items():
        data    = phase4_results.get(attr, {})
        content = data.get("content") or "No information found."
        sources = data.get("sources", [])

        source_tuples = []
        for s in sources:
            if isinstance(s, dict):
                source_tuples.append((s.get("url", ""), ""))
            elif isinstance(s, str) and s:
                source_tuples.append((s, ""))

        fields[placeholder] = [(content, source_tuples)]

    # NAME field
    fields["NAME"] = [(name, [])]

    return fields


def run(phase4_output: dict) -> str:
    name    = phase4_output.get("name", "Unknown")
    results = phase4_output.get("results", {})

    print(f"\n  Loading organisation template...")
    doc = Document(str(TEMPLATE_PATH))

    all_fields = format_org_fields(results, name)

    print(f"  Filling {len(all_fields)} placeholders...")
    populate_document_placeholders_with_dates(doc, all_fields)

    # Save to outputs/<OrgName>/
    safe_name  = re.sub(r"[^\w\s-]", "", name).strip().replace(" ", "_")
    person_dir = OUTPUTS_DIR / safe_name

    if person_dir.exists():
        import time
        for attempt in range(5):
            try:
                shutil.rmtree(person_dir, ignore_errors=True)
                if not person_dir.exists():
                    break
            except Exception:
                pass
            time.sleep(1)

    person_dir.mkdir(parents=True, exist_ok=True)
    output_path = str(person_dir / f"{safe_name}.docx")
    doc.save(output_path)

    print(f"  ✅ Document saved: {output_path}")
    return output_path
